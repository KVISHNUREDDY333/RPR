import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_review_report(report_data: dict, output_path: str) -> str:
    doc = Document()

    # Title
    title = doc.add_heading("AI Research Paper Review Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Overall Score
    doc.add_heading("Overall Score", level=1)
    score_para = doc.add_paragraph()
    run = score_para.add_run(f"{report_data.get('overall_score', 'N/A')} / 10")
    run.bold = True
    run.font.size = Pt(18)

    # Summary
    summary = report_data.get("summary", "")
    if summary:
        doc.add_heading("Executive Summary", level=1)
        doc.add_paragraph(summary)

    # Strengths
    strengths = report_data.get("strengths", [])
    if strengths:
        doc.add_heading("Strengths", level=1)
        for s in strengths:
            doc.add_paragraph(s, style="List Bullet")

    # Weaknesses
    weaknesses = report_data.get("weaknesses", [])
    if weaknesses:
        doc.add_heading("Areas for Improvement", level=1)
        for w in weaknesses:
            doc.add_paragraph(w, style="List Bullet")

    # Suggestions
    suggestions = report_data.get("suggestions", [])
    if suggestions:
        doc.add_heading("Recommendations", level=1)
        for s in suggestions:
            doc.add_paragraph(s, style="List Number")

    # Section Analysis
    section_analysis = report_data.get("section_analysis", [])
    if section_analysis:
        doc.add_heading("Section-by-Section Analysis", level=1)
        for section in section_analysis:
            doc.add_heading(section.get("title", "Section"), level=2)
            p = doc.add_paragraph()
            p.add_run(f"Score: {section.get('score', 'N/A')}/10  |  Clarity: {section.get('clarity', 'N/A')}").bold = True
            doc.add_paragraph(section.get("reason", ""))
            issues = section.get("issues", [])
            if issues:
                doc.add_paragraph("Issues identified:", style="List Bullet")
                for issue in issues:
                    doc.add_paragraph(issue, style="List Bullet 2")

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path
