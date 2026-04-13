from typing import List
from docx import Document
import os

def reconstruct(sections: List[dict], output_path: str) -> str:
    """Assemble refined sections into a DOCX file."""
    doc = Document()
    doc.add_heading("Refined Research Paper", 0)
    
    for section in sections:
        doc.add_heading(section["title"], level=1)
        for para in section["refined"].split("\n\n"):
            if para.strip():
                doc.add_paragraph(para.strip())
    
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path
